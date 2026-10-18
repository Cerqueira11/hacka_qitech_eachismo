import json
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


def run_workbook_path(report: dict, directory: Path) -> Path:
    timestamp = datetime.now().astimezone().strftime("%Y%m%d_%H%M%S")
    run_id = report.get("run_id", "run")
    return Path(directory) / f"dict_random_results_{timestamp}_{run_id}.xlsx"


def _response_data(scenario: dict) -> dict:
    result = scenario.get("post_result", {})
    response = result.get("response")
    if not isinstance(response, dict):
        return {}
    return response


def _dict_summary(query: dict) -> tuple:
    response = query.get("response", {}).get("ListFraudMarkersResponse", {})
    markers = response.get("FraudMarkers", {}).get("FraudMarker", [])
    marker_text = "; ".join(
        f"{marker.get('FraudType', '')}/{marker.get('Status', '')}"
        for marker in markers
    ) or "Sem marcações retornadas"
    return (
        response.get("Total", 0),
        response.get("HasMoreElements", False),
        marker_text,
    )


def _json_text(value) -> str:
    return json.dumps(value, ensure_ascii=False, separators=(",", ":")) if value is not None else ""


def create_results_workbook(report: dict, output_path: Path) -> Path:
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    workbook = Workbook()
    summary = workbook.active
    summary.title = "Resumo"
    transactions = workbook.create_sheet("Transações")
    markers_sheet = workbook.create_sheet("DICT")

    post_results = [
        scenario for scenario in report.get("scenarios", [])
        if "post_result" in scenario
    ]
    successful = [
        scenario for scenario in post_results
        if 200 <= scenario["post_result"].get("http_status", 0) < 300
    ]
    verdict_counts = {}
    for scenario in post_results:
        verdict = _response_data(scenario).get("status", "Sem veredito")
        verdict_counts[verdict] = verdict_counts.get(verdict, 0) + 1

    summary_rows = [
        ("Relatório gerado em", report.get("generated_at", "")),
        ("Seed", report.get("seed", "")),
        ("ID da execução", report.get("run_id", "")),
        ("Cenários preparados", report.get("scenario_count", len(report.get("scenarios", [])))),
        ("POSTs registrados no log", len(post_results)),
        ("POSTs HTTP 2xx", len(successful)),
        ("POSTs sem resultado registrado", len(report.get("scenarios", [])) - len(post_results)),
        ("Observação", (
            "O log não contém respostas POST; resultados do modelo não foram registrados."
            if not post_results else
            "Os vereditos abaixo são respostas reais da API registradas pelo runner."
        )),
    ]
    summary.append(["Indicador", "Valor"])
    for row in summary_rows:
        summary.append(list(row))
    summary.append([])
    summary.append(["Veredito da API", "Quantidade"])
    for verdict, count in sorted(verdict_counts.items()):
        summary.append([verdict, count])

    transaction_headers = [
        "Cenário", "ID transação", "Data/hora", "Valor (centavos)",
        "CPF/CNPJ origem", "Conta origem", "CPF/CNPJ destino", "Conta destino",
        "Dispositivo", "Canal", "IP", "HTTP POST",
        "Veredito", "Risk score", "Motivos",
        "Resposta completa da API",
    ]
    transactions.append(transaction_headers)

    marker_headers = [
        "Cenário", "Lado", "CPF/CNPJ consultado", "Total DICT",
        "HasMoreElements", "Marcações retornadas", "Resposta DICT completa",
    ]
    markers_sheet.append(marker_headers)

    for scenario in report.get("scenarios", []):
        payload = scenario.get("post_payload", {})
        source_account = payload.get("source_account", {})
        destination_account = payload.get("destination_account", {})
        source = payload.get("source", {})
        api_response = _response_data(scenario)
        metadata = api_response.get("metadata", {})
        post_result = scenario.get("post_result", {})
        error = post_result.get("error", "")

        transactions.append([
            scenario.get("fixture", ""),
            payload.get("id", ""),
            payload.get("transaction_date", ""),
            payload.get("amount", ""),
            source_account.get("owner", {}).get("document_number", ""),
            f"{source_account.get('participant', '')}/{source_account.get('branch', '')}/{source_account.get('account_number', '')}",
            destination_account.get("owner", {}).get("document_number", ""),
            f"{destination_account.get('participant', '')}/{destination_account.get('branch', '')}/{destination_account.get('account_number', '')}",
            source.get("device_id", ""),
            source.get("channel", ""),
            source.get("ip", ""),
            post_result.get("http_status", "Não registrado"),
            api_response.get("status", "Não registrado"),
            metadata.get("risk_score", ""),
            "; ".join(metadata.get("reasons", [])),
            _json_text(api_response) if api_response else error,
        ])

        for side, query in scenario.get("dict_queries", {}).items():
            total, has_more, markers = _dict_summary(query)
            markers_sheet.append([
                scenario.get("fixture", ""),
                side,
                query.get("document_number", ""),
                total,
                has_more,
                markers,
                _json_text(query.get("response", {})),
            ])

    header_fill = PatternFill("solid", fgColor="17365D")
    header_font = Font(color="FFFFFF", bold=True)
    for sheet in (summary, transactions, markers_sheet):
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions if sheet.max_row > 1 else "A1"
        for cell in sheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        sheet.row_dimensions[1].height = 30
        for column in sheet.columns:
            max_length = max(
                (len(str(cell.value)) for cell in column if cell.value is not None),
                default=10,
            )
            sheet.column_dimensions[get_column_letter(column[0].column)].width = min(
                max(max_length + 2, 12), 55
            )
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)

    transactions.column_dimensions["P"].width = 55
    markers_sheet.column_dimensions["G"].width = 55
    workbook.save(output_path)
    return output_path


def create_workbook_from_log(log_path: Path, output_path: Path) -> Path:
    with Path(log_path).open(encoding="utf-8") as log_file:
        report = json.load(log_file)
    return create_results_workbook(report, output_path)
